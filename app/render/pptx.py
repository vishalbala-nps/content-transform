"""Slide decks with python-pptx. No model calls.

A builder with one method per kind of slide, so the renderer knows layouts,
not any format's payload. Colours, font, organisation name and logo come
from a Theme (render/theme.py), the same one the PDFs use.

Every slide's title is a real title placeholder, restyled, rather than a
free text box: PowerPoint's outline view, slide navigator and screen readers
read slide titles from there.
"""

from io import BytesIO

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.parts.image import Image
from pptx.util import Emu, Inches, Pt

from app.render.theme import Theme

WHITE = RGBColor(0xFF, 0xFF, 0xFF)

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)  # 16:9
MARGIN = Inches(0.75)
CONTENT_WIDTH = WIDTH - 2 * MARGIN

# Default template layouts.
TITLE_LAYOUT, TITLE_ONLY_LAYOUT = 0, 5

FOOTER_CHARS = 90


def _colour(hex_colour: str) -> RGBColor:
    return RGBColor.from_string(hex_colour.lstrip("#").upper())


def _style(paragraph, size: int, color: RGBColor, font: str, bold: bool = False) -> None:
    paragraph.alignment = PP_ALIGN.LEFT
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def _frame(shape, anchor=MSO_ANCHOR.TOP, margin: int = 0):
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = margin
    return frame


def _block(slide, left, top, width, height, color: RGBColor):
    """A flat filled rectangle: no outline, and no shadow from the theme."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _bullet(paragraph, color: RGBColor, font: str) -> None:
    """A coloured bullet with a hanging indent. python-pptx has no API for bullets."""
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(Inches(0.35)))
    ppr.set("indent", str(-Inches(0.35)))
    # Schema order: after spacing (already set), before any run defaults.
    etree.SubElement(etree.SubElement(ppr, qn("a:buClr")), qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(ppr, qn("a:buSzPct")).set("val", "100000")
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", font)
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


class DeckBuilder:
    def __init__(self, title: str, footer: str, theme: Theme):
        self._prs = Presentation()
        self._prs.slide_width, self._prs.slide_height = WIDTH, HEIGHT
        props = self._prs.core_properties
        props.title = title
        props.author = props.last_modified_by = theme.org_name
        self._footer = footer if len(footer) <= FOOTER_CHARS else footer[: FOOTER_CHARS - 1] + "…"
        self._theme = theme
        self._font = theme.deck_font
        self._ink, self._muted = _colour(theme.ink), _colour(theme.muted)
        self._accent, self._accent_soft = _colour(theme.accent), _colour(theme.accent_soft)
        self._rule, self._on_accent_muted = _colour(theme.rule), _colour(theme.on_accent_muted)

    def _style(self, paragraph, size: int, color: RGBColor, bold: bool = False) -> None:
        _style(paragraph, size, color, self._font, bold)

    def _logo_size(self, max_width, max_height) -> tuple[Emu, Emu]:
        """The theme's logo, as large as fits the box, keeping its proportions."""
        width_px, height_px = Image.from_blob(self._theme.logo).size
        scale = min(max_width / width_px, max_height / height_px)
        return Emu(int(width_px * scale)), Emu(int(height_px * scale))

    def _logo(self, slide, left, top, width, height) -> None:
        slide.shapes.add_picture(BytesIO(self._theme.logo), left, top, width, height)

    def _slide(self, layout: int, notes: str):
        slide = self._prs.slides.add_slide(self._prs.slide_layouts[layout])
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return slide

    def _title(self, slide, text: str, top, height, size: int, color: RGBColor, anchor) -> None:
        title = slide.shapes.title
        title.left, title.top, title.width, title.height = MARGIN, top, CONTENT_WIDTH, height
        frame = _frame(title, anchor)
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE  # PowerPoint shrinks an overlong title
        frame.text = text
        self._style(frame.paragraphs[0], size, color, bold=True)
        frame.paragraphs[0].line_spacing = 1.05

    def _chrome(self, slide) -> None:
        """Accent mark above the title (and the logo, top right), then a rule,
        the deck title and the slide number at the foot."""
        _block(slide, MARGIN, Inches(0.55), Inches(0.6), Inches(0.07), self._accent)
        if self._theme.logo:
            width, height = self._logo_size(Inches(1.6), Inches(0.45))
            self._logo(slide, WIDTH - MARGIN - width, Inches(0.3), width, height)

        _block(slide, MARGIN, Inches(6.8), CONTENT_WIDTH, Pt(0.75), self._rule)

        number = len(self._prs.slides)
        for text, left, width, align in (
            (self._footer, MARGIN, CONTENT_WIDTH - Inches(1), PP_ALIGN.LEFT),
            (str(number), WIDTH - MARGIN - Inches(1), Inches(1), PP_ALIGN.RIGHT),
        ):
            box = slide.shapes.add_textbox(left, Inches(6.9), width, Inches(0.3))
            frame = _frame(box)
            frame.text = text
            self._style(frame.paragraphs[0], 10, self._muted)
            frame.paragraphs[0].alignment = align

    def title_slide(self, title: str, subtitle: str, notes: str) -> None:
        slide = self._slide(TITLE_LAYOUT, notes)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = self._accent
        self._title(slide, title, Inches(1.4), Inches(2.9), 40, WHITE, MSO_ANCHOR.BOTTOM)

        sub = slide.placeholders[1]
        sub.left, sub.top, sub.width, sub.height = MARGIN, Inches(4.55), CONTENT_WIDTH, Inches(1.2)
        frame = _frame(sub)
        frame.text = subtitle
        self._style(frame.paragraphs[0], 18, self._on_accent_muted)

        brand = slide.shapes.add_textbox(MARGIN, Inches(6.75), CONTENT_WIDTH, Inches(0.35))
        frame = _frame(brand)
        frame.text = self._theme.org_name
        self._style(frame.paragraphs[0], 11, self._on_accent_muted, bold=True)

        if self._theme.logo:
            # On a white panel: a logo is usually drawn for a light background,
            # and the accent behind it is the kit's, not the logo's.
            pad = Inches(0.15)
            width, height = self._logo_size(Inches(2.2), Inches(0.6))
            _block(slide, MARGIN, Inches(0.5), width + 2 * pad, height + 2 * pad, WHITE)
            self._logo(slide, MARGIN + pad, Inches(0.5) + pad, width, height)

    def bullets_slide(self, title: str, bullets: list[str], notes: str) -> None:
        slide = self._slide(TITLE_ONLY_LAYOUT, notes)
        self._title(slide, title, Inches(0.8), Inches(1.35), 28, self._ink, MSO_ANCHOR.TOP)

        body = slide.shapes.add_textbox(MARGIN, Inches(2.4), CONTENT_WIDTH, Inches(4.1))
        frame = _frame(body)
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for i, text in enumerate(bullets):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.line_spacing = 1.1
            paragraph.space_after = Pt(14)
            self._style(paragraph, 20, self._ink)
            _bullet(paragraph, self._accent, self._font)
        self._chrome(slide)

    def numbers_slide(self, title: str, numbers: list[tuple[str, str]], notes: str) -> None:
        """Up to four figures, each a large value over its label on a tinted card."""
        slide = self._slide(TITLE_ONLY_LAYOUT, notes)
        self._title(slide, title, Inches(0.8), Inches(1.35), 28, self._ink, MSO_ANCHOR.TOP)

        gap, padding = Inches(0.3), Inches(0.35)
        card_width = Emu(int((CONTENT_WIDTH - gap * (len(numbers) - 1)) / len(numbers)))
        # One size for every value, small enough that the longest fits on one
        # line: Arial bold averages about 0.6 em per character.
        longest = max(len(value) for value, _ in numbers)
        value_size = max(24, min(44, int(Emu(card_width - 2 * padding).pt / (0.6 * longest))))
        for i, (value, label) in enumerate(numbers):
            left = MARGIN + i * (card_width + gap)
            card = _block(slide, left, Inches(2.5), card_width, Inches(3.6), self._accent_soft)
            _block(slide, left, Inches(2.5), Inches(0.07), Inches(3.6), self._accent)

            frame = _frame(card, margin=padding)
            frame.text = value
            self._style(frame.paragraphs[0], value_size, self._accent, bold=True)
            frame.paragraphs[0].space_after = Pt(10)
            paragraph = frame.add_paragraph()
            paragraph.text = label
            self._style(paragraph, 16, self._ink)
        self._chrome(slide)

    def to_bytes(self) -> bytes:
        out = BytesIO()
        self._prs.save(out)
        return out.getvalue()
